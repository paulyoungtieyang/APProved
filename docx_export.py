"""
Step 5: "Agent writes the report in Word format."

Plain python-docx, no external dependencies (no LibreOffice, no Word
installation needed) so this runs anywhere. If a brand template is supplied
(params.brand_template), we open that as the base document instead of a
blank one, so headers/footers/styles carry over -- swap the
`Document(template_path)` line if your template lives somewhere else.
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import ClientData, SubmissionParams
from drafting import DraftSection


def export_docx(sections: list[DraftSection], client_data: ClientData,
                 params: SubmissionParams, out_path: str, status: str = "DRAFT") -> str:
    template = params.brand_template
    doc = Document(template) if template and Path(template).exists() else Document()

    title = doc.add_heading(f"{client_data.device_name} — EU MDR Technical Documentation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Manufacturer: {client_data.manufacturer}  |  "
        f"Regulation: {params.regulation}  |  Device class: {params.device_class}  |  "
        f"Status: {status}  |  Generated: {date.today().isoformat()}"
    )
    run.italic = True
    run.font.size = Pt(10)

    doc.add_page_break()

    for section in sections:
        doc.add_heading(section.title, level=1)
        for paragraph in section.text.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
